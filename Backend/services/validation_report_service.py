"""Automated PBI validation DOCX report built from persisted artifacts.

This module is intentionally isolated from the browser pipeline:

* It never launches Playwright.
* It never calls an LLM or vision model.
* It never contacts Power BI or triggers another export/DOM extraction.
* It only reads the persisted artifacts (validation_capture.json plus
  browser_metrics.json under ``output/reports/<run_id>/``) and the
  screenshot/export files those artifacts reference.

The generated report is a filled copy of the team's manual template
``PBI Report Validation - Template.docx``. The original template file is
never modified; each generation works from a python-docx in-memory copy and
preserves the template's headings, table structure, and layout.

Entry point:

    from services.validation_report_service import build_validation_report

    path = await build_validation_report(run_id)
"""

from __future__ import annotations

import asyncio
import copy
import logging
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from services.run_artifact_store import (
    load_browser_metrics,
    load_validation_capture,
)
from utils.config import (
    REPORT_DIR,
    SCREENSHOT_DIR,
    VALIDATION_REPORT_TEMPLATE,
)


logger = logging.getLogger(__name__)

# Fields that the automation cannot determine are intentionally left blank so
# the team can fill them in manually. The workspace NAME is the exception: a
# Power BI workspace id (the ``/groups/{guid}`` URL fragment) is never a name,
# so when no actual workspace name is available the field states so instead of
# presenting a derived id as a name.
_NOT_AVAILABLE = ""
_NOT_AVAILABLE_WORKSPACE = "Not available"

# Generic dictionary-key markers for a persisted refresh timestamp. Key-driven
# only - no dashboard location, coordinate, label, or date format is assumed.
_PERSISTED_REFRESH_KEYWORDS = (
    "refresh",
    "last_updated",
    "updated_at",
    "updated_as_of",
    "data_as_of",
)

_MATCH_STATUSES = ("match", "table_matched")

_CHECKS = ("kpis", "visuals", "filters", "buttons")

# Generic statement describing what a successfully applied filter is
# expected to do. Filled only for real, runtime-applied filter scenarios
# (never for DOM-only comparison rows, and never for results/values that
# were not actually observed).
_FILTER_EXPECTED_BEHAVIOR = (
    "Dashboard visuals update according to the selected filter"
)


class ValidationReportNotFound(Exception):
    """Raised when no persisted capture exists for a run_id."""


class ValidationReportError(Exception):
    """Raised when the DOCX report cannot be generated."""


def _resolve_template_path() -> Path:
    candidates = [
        Path(VALIDATION_REPORT_TEMPLATE),
        Path.home() / "excel_exports" / "PBI Report Validation - Template.docx",
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise ValidationReportError(
        "PBI validation report template not found; searched: "
        + "; ".join(str(path) for path in candidates)
    )


def _safe(value, default: str = _NOT_AVAILABLE) -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def _safe_filename_dashboard(name: str) -> str:
    """Turn a captured dashboard/report title into a filesystem-safe name.

    Matches the requested ``<Dashboard_Name>_<run_id>.docx`` convention
    (spaces/punctuation become underscores). Never derived from run_id and
    never hardcoded.
    """
    text = _safe(name, "Dashboard")
    text = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", text)
    text = re.sub(r"\s+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.strip(" ._") or "Dashboard"


def report_filename(run_id: str, dashboard_name: str) -> str:
    """Return the DOCX file name: ``<Dashboard_Name>_<run_id>.docx``.

    The dashboard name is the actual captured dashboard/report title (the
    same value used in the report's "Report Name" field); only its
    filesystem-safe form is embedded in the file name.
    """
    fragment = _safe_filename_dashboard(_safe(dashboard_name, "Dashboard"))
    return f"{fragment}_{run_id}.docx"


def _workspace_name(dashboard: dict | None) -> str:
    """Return an actual workspace name when the dashboard object already
    carries one, otherwise the honest "Not available" marker.

    The URL's ``/groups/{guid}`` fragment is a workspace *id*, never a
    workspace *name*, so it is never surfaced as a name. No value is derived
    from run_id.
    """
    return _safe((dashboard or {}).get("workspace_name"), _NOT_AVAILABLE_WORKSPACE)


def _resolve_screenshot_path(recorded: str | None) -> Path | None:
    """Return an existing screenshot file for a recorded path.

    Handles stale absolute paths (e.g. captures taken before the backend
    folder reorganisation) by re-rooting them under PROJECT_ROOT and, as a
    final fallback, searching the screenshots store by filename.
    """
    if not recorded:
        return None

    candidate = Path(recorded)
    if candidate.exists():
        return candidate

    rerooted = Path(
        str(recorded)
        .replace("\\output\\", "\\Backend\\output\\")
        .replace("/output/", "/Backend/output/")
    )
    if rerooted.exists():
        return rerooted

    matches = sorted(Path(SCREENSHOT_DIR).rglob(candidate.name))
    if matches:
        return matches[0]

    return None


def _scan_refresh(obj) -> str | None:
    """Return the first non-empty value stored under a refresh-ish key.

    Key-driven only (e.g. ``refresh_timestamp``, ``last_updated``,
    ``updated_at``) so no dashboard location, label, or date format is
    hardcoded. Handles both persisted dicts and flat key/value entries.
    """
    if isinstance(obj, dict):
        for key, value in obj.items():
            lowered = str(key).casefold()
            if any(
                token in lowered for token in _PERSISTED_REFRESH_KEYWORDS
            ) and isinstance(value, str) and value.strip():
                return value.strip()
        for value in obj.values():
            found = _scan_refresh(value)
            if found:
                return found
    elif isinstance(obj, list):
        for item in obj:
            found = _scan_refresh(item)
            if found:
                return found
    return None


def _extract_refresh_stamp(executions: list) -> str:
    for execution in executions or []:
        containers = (
            execution,
            execution.get("dashboard") or {},
            execution.get("metrics") or {},
            execution.get("extraction") or {},
            execution.get("visual_data") or {},
        )
        for container in containers:
            found = _scan_refresh(container)
            if found:
                return found
    return _NOT_AVAILABLE


def _report_display_name(source_name: str, target_name: str, metrics: list) -> str:
    candidates: dict[str, int] = {}
    for metric in metrics or []:
        if not isinstance(metric, dict):
            continue
        title = _safe(metric.get("page_title"), "")
        parts = title.split(" - ")
        if not parts:
            continue
        trailing = parts[-1].strip().casefold()
        if trailing not in ("power bi", "microsoft power bi"):
            continue
        # Power BI tab titles are "<Page> - <Report> - Power BI" for
        # multi-page reports and "<Report> - Power BI" for single-page
        # reports; the report/dashboard title is the segment directly before
        # the Power BI suffix. The most frequently observed segment across
        # all pages/sides wins, so a report name repeated on every page
        # beats any single page's title.
        if len(parts) >= 3:
            segment = parts[-2].strip()
        elif len(parts) == 2:
            segment = parts[0].strip()
        else:
            continue
        if segment:
            candidates[segment] = candidates.get(segment, 0) + 1
    if candidates:
        return max(candidates, key=candidates.get)
    if source_name and target_name and source_name != target_name:
        return f"{source_name} vs {target_name}"
    return source_name or _NOT_AVAILABLE


def resolve_dashboard_name(source_name: str, target_name: str, metrics: list) -> str:
    """Public resolver for the report/dashboard display name.

    Single sourced place so the API response (``dashboard_name``) and the
    DOCX report agree. Kept DOM-derived for now; the AI pipeline surfaces a
    ``dashboard_name`` when the hosted service returns one.
    """
    return _report_display_name(source_name, target_name, metrics)


@dataclass
class PagePair:
    page_name: str
    source: dict
    target: dict
    comparison: dict | None = None
    source_image: str | None = None
    target_image: str | None = None


@dataclass
class ReportContext:
    source_name: str = _NOT_AVAILABLE
    target_name: str = _NOT_AVAILABLE
    report_name: str = _NOT_AVAILABLE
    source_workspace_name: str = _NOT_AVAILABLE_WORKSPACE
    target_workspace_name: str = _NOT_AVAILABLE_WORKSPACE
    validation_date: str = _NOT_AVAILABLE
    source_refresh: str = _NOT_AVAILABLE
    target_refresh: str = _NOT_AVAILABLE
    pages: list[PagePair] = field(default_factory=list)
    filter_rows: list[list[str]] = field(default_factory=list)
    visual_rows: list[list[str]] = field(default_factory=list)
    export_rows: list[list[str]] = field(default_factory=list)
    overall_status: str = _NOT_AVAILABLE
    checks: int = 0
    mismatches: int = 0
    overall_match_percentage: float | None = None
    notes: list[str] = field(default_factory=list)
    ai_analysis: dict | None = None


def _pass_fail(status, extra: str = "") -> str:
    status = _safe(status, _NOT_AVAILABLE)
    lowered = status.casefold()
    if lowered in _MATCH_STATUSES:
        return "Pass"
    if lowered in (
        "mismatch",
        "missing in source",
        "missing in target",
        "table_mismatched",
        "table_mismatch",
        "table_not_compared",
        "not compared",
    ):
        suffix = f" - {extra}" if extra else ""
        return f"Fail{suffix}"
    if lowered in ("needs_review", "needs review"):
        return "Needs Review"
    return status


def _result_detail(item: dict, *, kind: str) -> str:
    status = _safe(item.get("status"))
    if status.casefold() != "match":
        if kind == "filters":
            return " | ".join(
                part
                for part in (
                    status,
                    f"Source: {_safe(item.get('source_selected'))}",
                    f"Target: {_safe(item.get('target_selected'))}",
                )
                if part != status
            )
        if kind == "buttons":
            return " | ".join(
                part
                for part in (
                    status,
                    f"Source selected: {_safe(item.get('source_selected'))}",
                    f"Target selected: {_safe(item.get('target_selected'))}",
                )
                if part != status
            )
        if kind == "kpis":
            return " | ".join(
                part
                for part in (
                    status,
                    f"Source: {_safe(item.get('source'))}",
                    f"Target: {_safe(item.get('target'))}",
                )
                if part != status
            )
        if kind == "visuals":
            return " | ".join(
                part
                for part in (
                    status,
                    f"Source: {_safe(item.get('source'))}",
                    f"Target: {_safe(item.get('target'))}",
                )
                if part != status
            )
    return status


def _build_context(document: dict, metrics_doc: dict, response: dict) -> ReportContext:
    ctx = ReportContext()
    ctx.validation_date = (document.get("created_at") or "")[:10] or _NOT_AVAILABLE

    groups = document.get("executions_by_dashboard") or []
    source_executions = groups[0] if len(groups) > 0 else []
    target_executions = groups[1] if len(groups) > 1 else []

    source = source_executions[0] if source_executions else {}
    target = target_executions[0] if target_executions else {}

    ctx.source_name = _safe((source.get("dashboard") or {}).get("name"), "Source")
    ctx.target_name = _safe((target.get("dashboard") or {}).get("name"), "Target")

    metrics = response.get("metrics") or []
    # The API response carries a resolved dashboard_name (an AI-provided
    # value when the hosted service returned one). Prefer it; fall back to
    # the deterministic DOM-derived name when absent.
    ctx.report_name = (
        (response.get("dashboard_name") or "").strip()
        or _report_display_name(
            ctx.source_name,
            ctx.target_name,
            metrics,
        )
    )

    ctx.source_workspace_name = _workspace_name(source.get("dashboard"))
    ctx.target_workspace_name = _workspace_name(target.get("dashboard"))
    ctx.source_refresh = _extract_refresh_stamp(source_executions)
    ctx.target_refresh = _extract_refresh_stamp(target_executions)

    comparisons_by_page = {
        item.get("page_name"): item
        for item in (response.get("comparison") or {}).get("page_comparisons", [])
    }

    for source_execution in source_executions:
        page_name = (source_execution.get("dashboard") or {}).get("page_name")
        if not page_name:
            continue
        target_execution = next(
            (
                item
                for item in target_executions
                if (item.get("dashboard") or {}).get("page_name") == page_name
            ),
            None,
        )
        if target_execution is None:
            continue
        ctx.pages.append(
            PagePair(
                page_name=page_name,
                source=source_execution,
                target=target_execution,
                comparison=comparisons_by_page.get(page_name),
                source_image=_resolve_screenshot_path(
                    (source_execution.get("metrics") or {}).get(
                        "screenshot_path"
                    )
                ),
                target_image=_resolve_screenshot_path(
                    (target_execution.get("metrics") or {}).get(
                        "screenshot_path"
                    )
                ),
            )
        )

    ctx.filter_rows = _build_filter_rows(document, ctx.pages)
    ctx.visual_rows, ctx.export_rows = _build_visual_and_export_rows(ctx.pages)

    ctx.checks, ctx.mismatches = _count_checks_and_mismatches(ctx.pages)
    ctx.overall_match_percentage = (
        (response.get("comparison") or {}).get("summary") or {}
    ).get("overall_match_percentage")
    ctx.overall_status = _derive_overall_status(
        response,
        len(ctx.pages),
        ctx.checks,
        ctx.mismatches,
    )

    baseline = document.get("baseline_screenshots") or {}
    for side in ("source", "target"):
        entry = baseline.get(side) or {}
        if entry.get("screenshot") and not entry.get("stable"):
            ctx.notes.append(
                f"{side.title()} baseline screenshot was skipped (dashboard "
                "did not reach a stable state)."
            )
    ctx.ai_analysis = document.get("ai_analysis") or response.get("ai_analysis")
    return ctx


def _seconds(value) -> str:
    if value is None:
        return _NOT_AVAILABLE
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return _NOT_AVAILABLE


def _build_filter_rows(document: dict, pages: list[PagePair]) -> list[list[str]]:
    rows: list[list[str]] = []
    row_id = 0
    scenario_by_page: dict[str, list[dict]] = {}
    for scenario in document.get("slicer_scenarios", []) or []:
        scenario_by_page.setdefault(
            scenario.get("page_name") or _scenario_page(scenario) or "", []
        ).append(scenario)

    def _emit_scenario(scenario: dict, default_page: str) -> None:
        nonlocal row_id
        row_id += 1
        applied = [scenario.get("source_applied"), scenario.get("target_applied")]
        passed = bool(applied[0] and applied[1])
        rows.append(
            [
                f"{row_id:02d}",
                scenario.get("page_name") or default_page,
                # Page/Tab Name and Filter Applied cover Page, Filter, and
                # Selected Value; the action column carries the actual
                # runtime-applied filter/value.
                f"{_safe(scenario.get('slicer'))} = "
                f"{_safe(scenario.get('value'))} | Filter applied on "
                f"Source & Target",
                _FILTER_EXPECTED_BEHAVIOR if passed else _NOT_AVAILABLE,
                "Pass" if passed else "Fail",
            ]
        )

    emitted: set[int] = set()
    for page in pages:
        page_name = page.page_name
        page_filter = (
            (page.source.get("dashboard") or {}).get("filter_applied")
            or "Default View"
        )
        comparison = page.comparison or {}
        for item in comparison.get("filters", []) or []:
            row_id += 1
            status = item.get("status")
            # DOM-only comparison rows: the expected-behavior column is
            # intentionally blank because no filter was applied here.
            rows.append(
                [
                    f"{row_id:02d}",
                    page_name,
                    f"Filter: {_safe(item.get('filter_name'))}",
                    _NOT_AVAILABLE,
                    _pass_fail(status, _result_detail(item, kind="filters")),
                ]
            )

        # Runtime filter scenarios, grouped by the page their slicers were
        # exercised on (single-page runs annotate their scenarios with a
        # page/tab name; any legacy scenario without one is resolved below).
        for scenario in scenario_by_page.get(page_name, []):
            if id(scenario) not in emitted:
                emitted.add(id(scenario))
                _emit_scenario(scenario, page_name)

    # Scenarios left over — those that carry no page/tab (legacy
    # single-page dicts) or whose page/tab did not match any captured
    # page — still earn a row so recorded results are never dropped.
    for scenarios in scenario_by_page.values():
        for scenario in scenarios:
            if id(scenario) in emitted:
                continue
            emitted.add(id(scenario))
            _emit_scenario(
                scenario,
                _scenario_page(scenario) or (pages[0].page_name if pages else "Unknown"),
            )

    if not rows:
        rows.append(
            [
                "01",
                "All",
                "No filter validation results recorded",
                _NOT_AVAILABLE,
                _NOT_AVAILABLE,
            ]
        )
    return rows


def _scenario_page(scenario: dict) -> str:
    """Best-effort page/tab name for a slicer scenario that does not carry
    one; falls back to an empty string rather than fabricating a value."""
    return _safe(scenario.get("page") or scenario.get("page_name"), "")


def _build_visual_and_export_rows(
    pages: list[PagePair],
) -> tuple[list[list[str]], list[list[str]]]:
    visual_rows: list[list[str]] = []
    export_rows: list[list[str]] = []
    visual_id = 0
    export_id = 0

    for page in pages:
        comparison = page.comparison or {}
        page_filter = (
            (page.source.get("dashboard") or {}).get("filter_applied")
            or "Default View"
        )

        def add_visual(text: str, status, detail: str = "") -> None:
            nonlocal visual_id
            visual_id += 1
            visual_rows.append(
                [
                    f"{visual_id:02d}",
                    page.page_name,
                    text,
                    f"Filter: {page_filter}",
                    _NOT_AVAILABLE,
                    _pass_fail(status, detail),
                ]
            )

        for kind, key, label in (
            ("kpis", "kpis", "KPI"),
            ("visuals", "visuals", "Visual"),
            ("buttons", "buttons", "Button group"),
        ):
            for item in comparison.get(key, []) or []:
                name = _safe(
                    item.get("kpi") or item.get("visual") or item.get("name")
                )
                add_visual(
                    f"{label}: {name}",
                    item.get("status"),
                    _result_detail(item, kind=kind),
                )

        tables = comparison.get("tables") or {}
        comparisons = (tables.get("comparisons") or []) or []
        for item in comparisons:
            name = (
                f"{_safe(item.get('source_table'))} vs "
                f"{_safe(item.get('target_table'))}"
            )
            add_visual(
                name,
                item.get("status"),
                _safe(item.get("reason")),
            )

        comparison_by_table: dict[str, dict] = {}
        for item in comparisons:
            for key in (
                _table_pairing_key(item),
                _normalise_table_title(item.get("source_table")),
                _normalise_table_title(item.get("target_table")),
            ):
                if key and key != "n/a":
                    comparison_by_table.setdefault(key, item)

        source_exports = (page.source.get("visual_data") or {}).get(
            "table_exports",
            [],
        ) or []
        target_exports = (page.target.get("visual_data") or {}).get(
            "table_exports",
            [],
        ) or []
        source_table_visuals = (
            (page.source.get("visual_data") or {}).get("table_visuals", []) or []
        )
        target_table_visuals = (
            (page.target.get("visual_data") or {}).get("table_visuals", []) or []
        )

        # Pair on _table_pairing_key, but keep the dashboard-specific display
        # title for the "Page / Table" column.
        display_title_by_key: dict[str, str] = {}
        dom_detected_keys: set[str] = set()
        for visual in [*source_table_visuals, *target_table_visuals]:
            key = _table_pairing_key(visual)
            if key:
                dom_detected_keys.add(key)
                display_title_by_key.setdefault(key, _safe(visual.get("title")))
        for export in [*source_exports, *target_exports]:
            key = _table_pairing_key(export)
            if key:
                display_title_by_key.setdefault(key, _safe(export.get("title")))

        dom_titles = sorted(display_title_by_key, key=str.casefold)

        if not dom_titles:
            export_id += 1
            export_rows.append(
                [
                    f"{export_id:02d}",
                    f"{page.page_name}: No table or matrix visual detected",
                    page_filter,
                    _NOT_AVAILABLE,
                    _NOT_AVAILABLE,
                    "Case D: No table/matrix visual detected on either report - "
                    "cannot compare exports",
                ]
            )
            continue

        for title in dom_titles:
            source_export = next(
                (
                    export
                    for export in source_exports
                    if _table_pairing_key(export) == title
                ),
                None,
            )
            target_export = next(
                (
                    export
                    for export in target_exports
                    if _table_pairing_key(export) == title
                ),
                None,
            )
            comparison_entry = comparison_by_table.get(title)
            dom_detected = title in dom_detected_keys
            export_id += 1
            name = _safe(
                (source_export or target_export or {}).get("title"),
            ) or display_title_by_key.get(title) or title
            export_rows.append(
                [
                    f"{export_id:02d}",
                    f"{page.page_name} / {name or title}",
                    page_filter,
                    _export_reference(source_export),
                    _export_reference(target_export),
                    _export_remarks(
                        source_export,
                        target_export,
                        comparison_entry,
                        dom_detected=dom_detected,
                    ),
                ]
            )

    return visual_rows, export_rows


def _normalise_table_title(title) -> str:
    return " ".join(str(title or "").casefold().split())


def _table_pairing_key(record: dict) -> str:
    """Identity used to pair a table across the two dashboards.

    A table the report gives no caption of its own is displayed as
    ``<dashboard>_table_<n>``, so its display name is dashboard-specific and
    cannot pair anything. ``comparison_key`` is the page-scoped ordinal that is
    identical on both dashboards, so it wins whenever present.
    """
    comparison_key = str((record or {}).get("comparison_key") or "").strip()
    if comparison_key:
        return comparison_key.casefold()
    return _normalise_table_title((record or {}).get("title"))


def _export_reference(export: dict | None) -> str:
    if not export:
        return _NOT_AVAILABLE
    if export.get("file_path"):
        reference = str(export["file_path"])
    elif export.get("status") in ("downloaded", "success"):
        reference = "Export succeeded (no file path recorded)"
    else:
        reference = f"Failed: {_safe(export.get('error'), 'no export produced')}"
    data = export.get("data")
    rows = None
    if isinstance(data, dict):
        rows = data.get("rows")
    if isinstance(rows, list):
        reference = f"{reference} ({len(rows)} rows)"
    return reference


def _is_browser_gone_error(error: str) -> bool:
    """True when a failure was caused by the browser closing, not by the visual.

    These read as "Export data is not available for this visual" in a report even
    though the export option was found and clicked, so they must be labelled
    separately or a working export looks permanently unsupported.
    """
    text = str(error or "").casefold()
    return (
        "has been closed" in text
        or "targetclosederror" in text
        or "target page closed" in text
        or "browser has been closed" in text
        or "connection closed" in text
    )


def _export_remarks(
    source_export: dict | None,
    target_export: dict | None,
    comparison: dict | None = None,
    dom_detected: bool = False,
) -> str:
    if not source_export and not target_export:
        if dom_detected:
            return (
                "Case E: Table/matrix visual detected on one or both reports but "
                "no export data was produced - Not compared"
            )
        return "Case D: No table/matrix export available"

    source_ok = bool(
        source_export
        and source_export.get("status") in ("downloaded", "success")
    )
    target_ok = bool(
        target_export
        and target_export.get("status") in ("downloaded", "success")
    )

    if not source_ok and not target_ok:
        # "Export data is not available for this visual" means Power BI's menu
        # genuinely offered no export for that visual. A browser that died
        # mid-save is a completely different fact and must not be reported as a
        # property of the visual, or a working export looks permanently broken.
        errors: list[str] = []
        unavailable: list[str] = []
        for export in (source_export, target_export):
            error = (export or {}).get("error")
            outcome = (export or {}).get("export_outcome")
            if outcome == "browser_closed" or (
                error and _is_browser_gone_error(error)
            ):
                errors.append(
                    _safe(error, "the browser closed before the export file was saved")
                )
            elif error:
                unavailable.append(_safe(error))
        if errors:
            return (
                "Case C: Not compared - the browser closed while saving the "
                f"export file (the visual does offer Export data): "
                f"{' | '.join(errors)}"
            )
        suffix = f": {' | '.join(unavailable)}" if unavailable else ""
        return f"Case C: Not compared - both exports failed{suffix}"
    if not (source_ok and target_ok):
        failed_side = "Target" if source_ok else "Source"
        failed_export = target_export if source_ok else source_export
        reason = _safe((failed_export or {}).get("error"), "no export produced")
        return f"Case B: Not compared - {failed_side} export failed: {reason}"
    if not comparison:
        return "Case E: Exports produced but no comparison result recorded"
    status = comparison.get("status")
    if status == "TABLE_MATCHED":
        return "Case A: Match (TABLES MATCHED)"
    if status != "TABLE_MISMATCHED":
        return (
            f"Case E: Not compared - status "
            f"{_safe(status, 'unreported')}"
        )
    parts = ["Case A: Mismatch",
             f"{_safe(comparison.get('source_row_count'))} source rows vs "
             f"{_safe(comparison.get('target_row_count'))} target rows"]
    matched = comparison.get("matched_row_count")
    if matched is not None:
        parts.append(f"{matched} matched rows")
    if comparison.get("reason"):
        parts.append(_safe(comparison.get("reason")))
    cells = comparison.get("cell_mismatches") or []
    if cells:
        parts.append(f"{len(cells)} differing cell values")
    missing_cols = comparison.get("missing_columns_in_target") or []
    extra_cols = comparison.get("extra_columns_in_target") or []
    if missing_cols:
        parts.append(f"columns only in source: {', '.join(map(str, missing_cols))}")
    if extra_cols:
        parts.append(f"columns only in target: {', '.join(map(str, extra_cols))}")
    missing_rows = comparison.get("missing_rows_in_target") or []
    extra_rows = comparison.get("extra_rows_in_target") or []
    if missing_rows:
        parts.append(f"{len(missing_rows)} rows only in source")
    if extra_rows:
        parts.append(f"{len(extra_rows)} rows only in target")
    return " | ".join(parts)


def _count_checks_and_mismatches(pages: list[PagePair]) -> tuple[int, int]:
    checks = 0
    mismatches = 0
    for page in pages:
        comparison = page.comparison or {}
        for key in _CHECKS:
            for item in comparison.get(key, []) or []:
                checks += 1
                if _safe(item.get("status")).casefold() != "match":
                    mismatches += 1
        tables = comparison.get("tables") or {}
        compared = int(tables.get("compared_table_count", 0) or 0)
        checks += compared
        mismatches += int(tables.get("mismatch_count", 0) or 0)
    return checks, mismatches


def _derive_overall_status(
    response: dict,
    page_count: int,
    checks: int,
    mismatches: int,
) -> str:
    comparison = response.get("comparison") or {}
    if comparison.get("status") != "success":
        status = _safe(comparison.get("status"), "not_compared")
        if status in ("success",):
            return _NOT_AVAILABLE
        return status.replace("_", " ").title() or _NOT_AVAILABLE
    if page_count == 0 or checks == 0:
        return "Not Compared"
    return "Pass" if mismatches == 0 else "Partial"


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _find_heading(document: Document, prefix: str):
    for paragraph in document.paragraphs:
        if prefix in paragraph.text and paragraph.style.name.startswith(
            "Heading"
        ):
            return paragraph
    return None


def _find_table(document: Document, *needles: str) -> object | None:
    wanted = [needle.casefold() for needle in needles if needle]
    for table in document.tables:
        text = " | ".join(
            cell.text.casefold().strip()
            for row in table.rows[:2]
            for cell in row.cells
        )
        if any(needle in text for needle in wanted):
            return table
    return None


def _clear_rows_after_table_header(table) -> None:
    for row in list(table.rows[1:]):
        row._tr.getparent().remove(row._tr)


def _append_rows(table, rows: list[list[str]], template=None) -> None:
    """Append rows without relying on ``add_row``.

    The team template stores fractional (non-integer twips) ``w:gridCol``
    widths, which python-docx's ``Table.add_row`` cannot parse. Appending a
    deep copy of an existing XML row with the correct per-cell layout avoids
    touching the grid entirely.
    """
    if template is None:
        template = copy.deepcopy(table.rows[-1]._tr) if table.rows else None
    for row_values in rows:
        if template is None:
            cells = table.add_row().cells
        else:
            table._tbl.append(copy.deepcopy(template))
            cells = table.rows[-1].cells
        for index, value in enumerate(row_values):
            if index < len(cells):
                cells[index].text = str(value)


def _fill_environment_table(
    table,
    ctx: ReportContext,
) -> None:
    # The template's own labels ("Snowflake Work Space Name" /
    # "GCP Work Space Name" / etc.) are preserved; only values are filled so
    # the source/target meaning stays attached to the correct row.
    rows = table.rows
    rows[0].cells[1].text = ctx.report_name
    rows[1].cells[1].text = _NOT_AVAILABLE
    rows[2].cells[1].text = ctx.source_workspace_name
    rows[3].cells[1].text = ctx.target_workspace_name
    rows[4].cells[1].text = ""
    rows[5].cells[1].text = ctx.validation_date
    rows[6].cells[1].text = ctx.overall_status

    _append_rows(
        table,
        [
            ["Report Generated", date.today().isoformat()],
        ],
    )


def _fill_refresh_table(table, ctx: ReportContext) -> None:
    table.rows[0].cells[1].text = ctx.source_name
    table.rows[0].cells[2].text = ctx.target_name
    table.rows[1].cells[1].text = ctx.source_refresh
    table.rows[1].cells[2].text = ctx.target_refresh
    table.rows[2].cells[1].text = ""
    table.rows[2].cells[2].text = ""
    table.rows[3].cells[1].text = ""
    table.rows[3].cells[2].text = ""
    table.rows[4].cells[1].text = ""
    table.rows[4].cells[2].text = ""


def _fill_summary_table(table, ctx: ReportContext) -> None:
    table.rows[0].cells[1].text = ctx.overall_status
    table.rows[1].cells[1].text = ""
    table.rows[2].cells[1].text = ""
    extra_rows = [
        ("Total Validations / Checks", str(ctx.checks)),
        ("Total Mismatches", str(ctx.mismatches)),
        (
            "Overall Match (%)",
            (
                f"{ctx.overall_match_percentage:.2f}"
                if ctx.overall_match_percentage is not None
                else _NOT_AVAILABLE
            ),
        ),
    ]
    for label, value in extra_rows:
        _append_rows(table, [[label, value]])


# US-Letter portrait with 1.0in margins leaves a 6.5in printable width.
# The reference report embeds full-width screenshots (6.5x~3.25) under
# captioned sub-steps (e.g. "2.1 GCP"); the automated report follows this
# with a full-width source/target pair per page, aspect preserved.
_SCREENSHOT_PAGE_WIDTH_IN = 6.5
_SCREENSHOT_PAGE_MAX_HEIGHT_IN = 9.0


def _fit_image_size(
    image_path: Path,
    max_width_in: float,
    max_height_in: float,
) -> tuple[Inches, Inches]:
    """Return an aspect-preserving image size that fits within the box.

    Falls back to max-width-only sizing when the pixel dimensions cannot be
    read (unknown format, unavailable PIL), which Word still renders without
    distortion because it derives the height from the width automatically.
    """
    try:
        from PIL import Image

        with Image.open(str(image_path)) as image:
            width_px, height_px = image.size
    except Exception:
        logger.debug(
            "Image dimensions unavailable; using width-only sizing | path=%s",
            image_path,
        )
        return Inches(max_width_in), None

    if width_px <= 0 or height_px <= 0:
        return Inches(max_width_in), None

    scale = min(
        max_width_in / width_px,
        max_height_in / height_px,
    )
    return Inches(width_px * scale), Inches(height_px * scale)


def _add_image_paragraph(
    anchor,
    image_path: Path,
    max_width_in: float = _SCREENSHOT_PAGE_WIDTH_IN,
    max_height_in: float = _SCREENSHOT_PAGE_MAX_HEIGHT_IN,
) -> None:
    """Embed one screenshot at full printable width, aspect preserved."""
    if not image_path or not Path(image_path).is_file():
        return
    try:
        paragraph = anchor.insert_paragraph_before()
        run = paragraph.add_run()
        width, height = _fit_image_size(
            Path(image_path),
            max_width_in,
            max_height_in,
        )
        run.add_picture(str(image_path), width=width, height=height)
    except Exception:
        logger.exception("Failed to embed screenshot | path=%s", image_path)


def _insert_screenshots_block(document: Document, ctx: ReportContext) -> None:
    section5 = _find_heading(document, "5. Data Export Validation Log")
    if section5 is None:
        logger.warning("Section 5 heading not found; screenshots skipped")
        return

    inserted_any = False
    for page in ctx.pages:
        if not page.source_image and not page.target_image:
            continue
        inserted_any = True

        summary = _page_summary_text(page)

        caption = section5.insert_paragraph_before(
            f"{page.page_name} - Page Screenshots (Source vs Target)"
            + (f" | {summary}" if summary else "")
        )
        for run in caption.runs:
            run.bold = True

        if page.source_image:
            section5.insert_paragraph_before(
                f"Source: {ctx.source_name}"
            )
            _add_image_paragraph(section5, Path(page.source_image))

        if page.target_image:
            section5.insert_paragraph_before(
                f"Target: {ctx.target_name}"
            )
            _add_image_paragraph(section5, Path(page.target_image))

    if not inserted_any:
        logger.warning("No screenshots found to embed in the report")


def _page_summary_text(page: PagePair) -> str:
    summary = (page.comparison or {}).get("summary") or {}
    percentage = summary.get("overall_match_percentage")
    if percentage is None:
        return "Not compared"
    return f"Overall match: {float(percentage):.2f}%"


def _validate_pages_structure(ctx: ReportContext) -> None:
    if not ctx.pages:
        ctx.notes.append("No matching pages found between the two dashboards.")


_AI_TABLE_FONT_PT = 8.5


def _ai_payload(ai: dict) -> dict:
    results = ai.get("results")
    if isinstance(results, dict):
        return results
    return ai.get("compare_status_payload") or {}


def _ai_summary_rows(ai: dict) -> list[list[str]]:
    payload = _ai_payload(ai)
    rows: list[list[str]] = []

    def add(label: str, value) -> None:
        if value in (None, "", _NOT_AVAILABLE):
            return
        rows.append([label, str(value)])

    add("Status", ai.get("status"))
    add("Reason / Error", ai.get("reason") or ai.get("error"))
    add("Job ID", ai.get("job_id"))
    uploaded = ai.get("uploaded") or {}
    if uploaded.get("source") is not None or uploaded.get("target") is not None:
        add(
            "Images uploaded (Source / Target)",
            f"{_safe(uploaded.get('source'), '0')} / "
            f"{_safe(uploaded.get('target'), '0')}",
        )
    add("Dashboard (per AI)", ai.get("dashboard_name"))
    add("Duration (seconds)", _seconds(ai.get("duration_seconds")))
    total = payload if ai.get("status") == "completed" else ai
    add("LLM calls", total.get("total_llm_calls"))
    add("Total tokens", total.get("total_tokens"))
    add("Total prompt tokens", total.get("total_prompt_tokens"))
    add("Total image tokens", total.get("total_image_tokens"))
    add("Total input tokens", total.get("total_input_tokens"))
    add("Total output tokens", total.get("total_output_tokens"))
    add("Total cost (USD)", total.get("total_cost_usd"))
    add("Total cost (INR)", total.get("total_cost_inr"))
    workbook = ai.get("workbook_available") or payload.get("workbook_available")
    add("Workbook available", "Yes" if workbook else None)
    add("Workbook download", ai.get("download_url") or payload.get("download_url"))
    add("JSON download", ai.get("json_download_url") or payload.get("json_download_url"))
    return rows


def _ai_pairs_rows(pairs: list) -> list[list[str]]:
    return [
        [
            _safe(pair.get("pair")),
            _safe(pair.get("spartnash_title")),
            _safe(pair.get("trendence_title")),
            _safe(pair.get("total_items")),
            _safe(pair.get("matches")),
            _safe(pair.get("differences")),
            _safe(pair.get("spartnash_only")),
            _safe(pair.get("trendence_only")),
            _safe(pair.get("uncertain")),
            (
                f"{float(pair.get('match_percentage')) * 100:.1f}%"
                if pair.get("match_percentage") is not None
                else _NOT_AVAILABLE
            ),
        ]
        for pair in pairs
    ]


def _ai_image_stats_rows(image_stats: list) -> list[list[str]]:
    return [
        [
            _safe(item.get("folder")),
            _safe(item.get("image")),
            (
                f"{item.get('width')}x{item.get('height')}"
                if item.get("width") is not None and item.get("height") is not None
                else _NOT_AVAILABLE
            ),
            _safe(item.get("prompt_tokens")),
            _safe(item.get("image_input_tokens")),
            _safe(item.get("total_input_tokens")),
            _safe(item.get("output_tokens")),
            _safe(item.get("total_tokens")),
            _safe(item.get("total_cost_usd")),
            _safe(item.get("total_cost_inr")),
            _safe(item.get("llm_calls")),
        ]
        for item in image_stats
    ]


def _set_grid_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "BFBFBF")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _move_before(block, anchor) -> None:
    if anchor is not None:
        anchor.addprevious(block._p if hasattr(block, "_p") else block._tbl)


def _render_ai_table(
    doc: Document,
    rows: list[list[str]],
    anchor,
    *,
    bold_header: bool = False,
) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    _set_grid_borders(table)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            if col_index >= len(table.rows[row_index].cells):
                continue
            cell = table.rows[row_index].cells[col_index]
            cell.text = str(value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(_AI_TABLE_FONT_PT)
                if bold_header and row_index == 0:
                    run.font.bold = True
    _move_before(table, anchor)


def _render_ai_section(doc: Document, ctx: ReportContext) -> None:
    ai = ctx.ai_analysis
    if not ai:
        return

    anchor = _find_heading(doc, "6. Summary")
    anchor_paragraph = anchor._p if anchor is not None else None

    heading = doc.add_paragraph("7. AI Validation Results (Gemini Backed)")
    try:
        heading.style = doc.styles["Heading 1"]
    except KeyError:
        pass
    _move_before(heading, anchor_paragraph)

    summary = _ai_summary_rows(ai)
    if summary:
        _render_ai_table(doc, summary, anchor_paragraph)

    payload = _ai_payload(ai)
    pairs = payload.get("pairs") if isinstance(payload.get("pairs"), list) else []
    if pairs:
        header = [
            "Pair",
            "Source Title",
            "Target Title",
            "Total Items",
            "Matches",
            "Differences",
            "Source-only",
            "Target-only",
            "Uncertain",
            "Match %",
        ]
        _render_ai_table(
            doc,
            [header] + _ai_pairs_rows(pairs),
            anchor_paragraph,
            bold_header=True,
        )

    image_stats = (
        payload.get("image_stats")
        if isinstance(payload.get("image_stats"), list)
        else []
    )
    if image_stats:
        header = [
            "Folder",
            "Image",
            "Dimensions",
            "Prompt Tokens",
            "Image Tokens",
            "Input Tokens",
            "Output Tokens",
            "Total Tokens",
            "Cost USD",
            "Cost INR",
            "LLM Calls",
        ]
        _render_ai_table(
            doc,
            [header] + _ai_image_stats_rows(image_stats),
            anchor_paragraph,
            bold_header=True,
        )

    logger.info(
        "AI validation section rendered | summary_rows=%d | pairs=%d | image_stats=%d",
        len(summary),
        len(pairs),
        len(image_stats),
    )


def _render(
    run_id: str,
    document: dict,
    metrics_doc: dict,
    response: dict,
    destination: Path | None,
) -> Path:
    template_path = _resolve_template_path()
    try:
        doc = Document(str(template_path))
    except Exception as exc:
        raise ValidationReportError(
            f"Failed to open report template: {template_path}"
        ) from exc

    ctx = _build_context(document, metrics_doc, response)
    _validate_pages_structure(ctx)

    title_paragraph = None
    for paragraph in doc.paragraphs[:8]:
        if paragraph.text.strip() and paragraph.text.strip() != "Report Name":
            title_paragraph = paragraph
            break
    if title_paragraph is not None:
        _set_paragraph_text(title_paragraph, ctx.report_name)

    environment_table = _find_table(doc, "Dashboard / Report Name")
    refresh_table = _find_table(doc, "Refresh Date & Time")
    filter_table = _find_table(doc, "Expected Behavior")
    visual_table = _find_table(doc, "Visual Name")
    export_table = _find_table(doc, "BigQuery Sheet Name")
    summary_table = _find_table(doc, "Sign-Off Status")

    if environment_table is not None:
        _fill_environment_table(environment_table, ctx)
    if refresh_table is not None:
        _fill_refresh_table(refresh_table, ctx)
    if filter_table is not None:
        template = (
            copy.deepcopy(filter_table.rows[1]._tr)
            if len(filter_table.rows) > 1
            else None
        )
        _clear_rows_after_table_header(filter_table)
        _append_rows(filter_table, ctx.filter_rows, template)
    if visual_table is not None:
        template = (
            copy.deepcopy(visual_table.rows[1]._tr)
            if len(visual_table.rows) > 1
            else None
        )
        _clear_rows_after_table_header(visual_table)
        _append_rows(visual_table, ctx.visual_rows, template)
    if export_table is not None:
        template = (
            copy.deepcopy(export_table.rows[1]._tr)
            if len(export_table.rows) > 1
            else None
        )
        export_table.rows[0].cells[3].text = "Source Export Reference"
        export_table.rows[0].cells[4].text = "Target Export Reference"
        _clear_rows_after_table_header(export_table)
        _append_rows(export_table, ctx.export_rows, template)
    if summary_table is not None:
        _fill_summary_table(summary_table, ctx)

    _insert_screenshots_block(doc, ctx)
    _render_ai_section(doc, ctx)

    if ctx.notes:
        for note in ctx.notes:
            summary_heading = _find_heading(doc, "6. Summary")
            if summary_heading is not None:
                summary_heading.insert_paragraph_before(f"Note: {note}")

    destination = destination or (
        REPORT_DIR
        / run_id
        / report_filename(run_id, ctx.report_name)
    )
    destination.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(destination))
    except Exception as exc:
        raise ValidationReportError(
            f"Failed to save DOCX report to {destination}"
        ) from exc
    logger.info(
        "PBI validation report generated | run_id=%s | path=%s",
        run_id,
        destination,
    )
    return destination


async def build_validation_report(
    run_id: str,
    *,
    destination: Path | None = None,
) -> Path:
    """Generate the automated PBI validation DOCX for a persisted run.

    Operates entirely from the persisted validation capture:
    ``load_validation_capture`` + the pure-compute validation response.
    No browser, no LLM, no Power BI export, no DOM extraction.

    Raises ``ValidationReportNotFound`` when no capture exists for run_id.
    """
    document = load_validation_capture(run_id)
    if document is None:
        raise ValidationReportNotFound(
            f"No validation capture found for run_id {run_id}."
        )
    metrics_doc = load_browser_metrics(run_id) or {}

    from orchestration.validator import DashboardValidator

    validator = DashboardValidator()
    response = await validator.run_validation_from_artifacts(run_id)
    if response is None:
        raise ValidationReportNotFound(
            f"No validation capture found for run_id {run_id}."
        )

    return await asyncio.to_thread(
        _render,
        run_id,
        document,
        metrics_doc,
        response,
        destination,
    )